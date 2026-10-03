import csv
import json
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
DOCX_PATH = ROOT / "docs" / "Doc2Patient_v1_PoC_Revised.docx"
LOAD_DIR = ROOT / "docs" / "load_tests"
AUDIT_DIR = ROOT / "docs" / "audits"
FIGURES_DIR = ROOT / "docs" / "figures"


def load_results():
    results = []
    for users in (10, 25, 50, 100):
        path = LOAD_DIR / f"locust-{users}-clean_stats.csv"
        with path.open(newline="", encoding="utf-8-sig") as handle:
            row = next(row for row in csv.DictReader(handle) if row["Type"] == "")
        results.append({
            "users": users,
            "requests": int(row["Request Count"]),
            "errors": int(row["Failure Count"]),
            "rps": float(row["Requests/s"]),
            "p50": float(row["50%"]),
            "p95": float(row["95%"]),
            "maximum": float(row["Max Response Time"]),
        })
    return results


def load_audits():
    audits = []
    for label, filename in (
        ("Landing", "lighthouse-landing.json"),
        ("Login", "lighthouse-login.json"),
        ("FAQ", "lighthouse-faq.json"),
    ):
        data = json.loads((AUDIT_DIR / filename).read_text(encoding="utf-8"))
        audits.append({
            "label": label,
            "performance": round(data["categories"]["performance"]["score"] * 100),
            "accessibility": round(data["categories"]["accessibility"]["score"] * 100),
            "lcp": data["audits"]["largest-contentful-paint"]["numericValue"] / 1000,
            "bytes": data["audits"]["total-byte-weight"]["numericValue"] / 1024,
            "requests": len(data["audits"]["network-requests"]["details"]["items"]),
        })
    return audits


def make_figures(load_results, audit_results):
    FIGURES_DIR.mkdir(exist_ok=True)
    users = [item["users"] for item in load_results]
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.8), constrained_layout=True)
    axes[0].plot(users, [item["p50"] for item in load_results], marker="o", label="p50")
    axes[0].plot(users, [item["p95"] for item in load_results], marker="o", label="p95")
    axes[0].set(title="Response time by concurrent users", xlabel="Concurrent users", ylabel="Response time (ms)")
    axes[0].grid(alpha=0.3)
    axes[0].legend()
    axes[1].plot(users, [item["rps"] for item in load_results], marker="o", color="#1479b8")
    axes[1].set(title="Throughput by concurrent users", xlabel="Concurrent users", ylabel="Requests per second")
    axes[1].grid(alpha=0.3)
    fig.suptitle("Doc2Patient v1 local Locust evaluation", fontsize=13)
    fig.savefig(FIGURES_DIR / "load_test_results.png", dpi=240, bbox_inches="tight")
    plt.close(fig)

    labels = [item["label"] for item in audit_results]
    positions = range(len(labels))
    width = 0.36
    fig, axis = plt.subplots(figsize=(7.5, 4.4), constrained_layout=True)
    axis.bar([p - width / 2 for p in positions], [item["performance"] for item in audit_results], width, label="Performance")
    axis.bar([p + width / 2 for p in positions], [item["accessibility"] for item in audit_results], width, label="Accessibility")
    axis.set_ylim(0, 100)
    axis.set_ylabel("Lighthouse score")
    axis.set_xticks(list(positions), labels)
    axis.set_title("Doc2Patient v1 Lighthouse audit scores")
    axis.grid(axis="y", alpha=0.3)
    axis.legend()
    fig.savefig(FIGURES_DIR / "audit_scores.png", dpi=240, bbox_inches="tight")
    plt.close(fig)


def replace_picture(paragraph, image_path, width):
    for picture in paragraph._p.xpath(".//pic:pic"):
        picture.getparent().getparent().remove(picture.getparent())
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def update_document(load_results, audit_results):
    document = Document(DOCX_PATH)
    paragraphs = document.paragraphs
    total_requests = sum(item["requests"] for item in load_results)
    total_errors = sum(item["errors"] for item in load_results)
    highest = load_results[-1]

    paragraphs[2].text = (
        "Abstract—Access to timely healthcare consultation remains a critical challenge in rural and underserved communities across Ghana. "
        "This paper presents Doc2Patient, a proof-of-concept web-based teleconsultation platform combining browser WebRTC consultation, "
        "persistent chat, appointment management, and patient-recorded vital-sign sharing. Version 1 provides role-separated patient, doctor, "
        "and administrator accounts, optional two-factor authentication, and an interface in English, French, and Twi. A local HTTP load "
        f"evaluation completed {total_requests} requests across 10–100 simulated users without recorded request failures, but latency increased "
        f"to a {highest['p50'] / 1000:.1f}-second median and {highest['p95'] / 1000:.1f}-second 95th percentile at 100 users. Lighthouse performance scores "
        "were 95–100 and accessibility scores were 100 on the tested landing, login, and FAQ pages. An unauthenticated OWASP ZAP baseline found "
        "0 High, 3 Medium, 0 Low, and 6 Informational alerts after security-header hardening. These results are local proof-of-concept evidence, "
        "not production scalability, Ghanaian 3G/4G performance, clinical utility, or complete security assurance."
    )
    paragraphs[53].text = (
        "Flask-Babel provides the interface in English, French, and Twi (Akan); Ga and Ewe are planned. The responsive interface uses a locally "
        "built and minified Tailwind CSS bundle. Automated Lighthouse accessibility scores were 100 on the landing, login, and FAQ pages, while "
        "axe found ten moderate region-landmark findings on the landing page and no violations on the login page. These automated results do not "
        "replace manual screen-reader, keyboard, and usability testing."
    )
    paragraphs[61].text = (
        "The evaluation used a local Windows 11/Python 3.14 environment with one Flask-SocketIO development process and SQLite. Locust stages ran "
        "for 30 seconds at 10, 25, 50, and 100 simulated users with a spawn rate of 10 users per second. The test exercised the public landing "
        "page, login, and authenticated role-dispatch route. Lighthouse 13.5 and axe-core 4.13 tested the current optimized frontend; OWASP ZAP "
        "2.17.0 performed an unauthenticated quick baseline scan against the local application."
    )
    paragraphs[62].text = (
        "Functional smoke testing verified the public landing page, login form, CSRF token presence, protected-route redirects, role-dispatch routing, "
        "and rejection of the obsolete /dashboard path. The load test measured HTTP page and authentication traffic only; it did not measure WebRTC "
        "media, STUN/TURN connectivity, vitals-sharing latency, or pilot usability."
    )
    paragraphs[63].text = (
        "The current unauthenticated ZAP baseline reported 0 High, 3 Medium, 0 Low, and 6 Informational alerts. The remaining Medium alerts concern "
        "CSP policy quality: wildcard WebSocket schemes and use of unsafe-inline for scripts and styles. Security headers, cookie SameSite settings, "
        "and Socket.IO SRI protection are present. This scan does not establish complete platform security; authenticated authorization, CSRF, XSS, "
        "and real-time-channel testing remain required."
    )
    paragraphs[65].text = (
        "Functional. Route smoke checks passed for the public and protected-route behaviours described above. Load. Across the four clean 30-second "
        f"Locust stages, the test completed {total_requests} requests with {total_errors} recorded failures. At 100 users, the median response time "
        f"was {highest['p50'] / 1000:.1f} seconds, the 95th percentile was {highest['p95'] / 1000:.1f} seconds, and the maximum was "
        f"{highest['maximum'] / 1000:.1f} seconds. Latency increased substantially as concurrency rose."
    )
    paragraphs[68].text = "Local test instance: Windows 11, Python 3.14, one Flask-SocketIO development process, SQLite, and 30-second stages. CPU and memory were not sampled with a dedicated monitoring tool."
    paragraphs[69].text = "Ghanaian 3G/4G profiles, packet loss, STUN/TURN connectivity, WebRTC media statistics, and vitals-sharing latency were not measured. Lighthouse results used local Chrome defaults and must not be interpreted as Ghanaian network measurements."
    paragraphs[72].text = "No constrained-network profiles were applied in this evaluation."
    audit_summary = "; ".join(f"{item['label']}: performance {item['performance']}, accessibility {item['accessibility']}" for item in audit_results)
    paragraphs[73].text = (
        f"Page weight and accessibility. Lighthouse results were {audit_summary}. Landing transferred approximately {audit_results[0]['bytes']:.0f} KiB in "
        f"{audit_results[0]['requests']} requests with LCP {audit_results[0]['lcp']:.2f} seconds; login transferred {audit_results[1]['bytes']:.0f} KiB in "
        f"{audit_results[1]['requests']} requests with LCP {audit_results[1]['lcp']:.2f} seconds; FAQ transferred {audit_results[2]['bytes']:.0f} KiB in "
        f"{audit_results[2]['requests']} requests with LCP {audit_results[2]['lcp']:.2f} seconds. These are local Chrome results, not Slow 3G/4G measurements. "
        "Security. The current unauthenticated ZAP baseline reported 0 High, 3 Medium, 0 Low, and 6 Informational alerts."
    )
    paragraphs[76].text = "Tool: OWASP ZAP 2.17.0; mode: unauthenticated quick baseline scan; target: local http://127.0.0.1:5000/; report: docs/audits/zap-baseline.html."
    paragraphs[77].text = "Pilot: no participant study was conducted for this revision; usability and SUS results are not reported."
    paragraphs[85].text = (
        "Doc2Patient v1 demonstrates that a proof-of-concept teleconsultation platform combining WebRTC video and audio, persistent chat, appointment management, "
        "patient-recorded vital signs, and a multilingual interface can be built with open-source components. The local load baseline completed all recorded "
        f"requests without failures, but latency increased substantially at 100 users. Lighthouse scores were high on the tested pages, while the ZAP baseline "
        "retains three CSP-quality findings. The results support iterative development, not production scalability, Ghanaian network performance, clinical readiness, or complete security assurance."
    )
    paragraphs[89].text = "No pilot participant study was conducted as part of this evaluation; ethics approval and consent details are therefore not applicable to the automated tests. Test accounts used simulated data only."
    paragraphs[91].text = "The source code for Doc2Patient v1 is publicly available at https://github.com/FinalYear4/Doc2Patient. These results were generated from the local code using SQLite and a single development process. It is a research prototype: do not enter real patient data."

    load_table = document.tables[4]
    for row, result in zip(load_table.rows[1:], load_results):
        row.cells[0].text = str(result["users"])
        row.cells[1].text = f"{result['rps']:.1f}"
        row.cells[2].text = f"{result['p50']:.0f}"
        row.cells[3].text = f"{result['p95']:.0f}"
        row.cells[4].text = f"{result['errors'] / result['requests'] * 100:.1f}"

    network_table = document.tables[5]
    for row in network_table.rows[1:]:
        for cell in row.cells[1:]:
            cell.text = "Not measured"

    security_table = document.tables[6]
    counts = [
        ("High", "0", "None"),
        ("Medium", "3", "CSP wildcard / unsafe-inline policy warnings"),
        ("Low", "0", "None"),
        ("Informational", "6", "ZAP informational observations"),
    ]
    for row, values in zip(security_table.rows[1:], counts):
        for cell, value in zip(row.cells, values):
            cell.text = value

    replace_picture(paragraphs[57], FIGURES_DIR / "load_test_results.png", 6.4)
    replace_picture(paragraphs[73], FIGURES_DIR / "audit_scores.png", 5.8)
    document.save(DOCX_PATH)


if __name__ == "__main__":
    load_results_data = load_results()
    audit_results_data = load_audits()
    make_figures(load_results_data, audit_results_data)
    update_document(load_results_data, audit_results_data)
    print(f"Updated {DOCX_PATH}")
